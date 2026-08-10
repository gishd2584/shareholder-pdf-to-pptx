#!/usr/bin/env python3
"""Create a native editable multi-level ownership chart PowerPoint."""

from __future__ import annotations

import argparse
import json
import math
from collections import defaultdict
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


BLACK = RGBColor(26, 26, 26)
WHITE = RGBColor(255, 255, 255)
BLUE = RGBColor(33, 150, 243)
DARK_BLUE = RGBColor(24, 118, 190)
ORANGE = RGBColor(245, 124, 0)
RED = RGBColor(239, 68, 68)
FONT = "Arial Unicode MS"


def number(value: Any, fallback: float) -> float:
    return float(value) if isinstance(value, (int, float)) and math.isfinite(value) else fallback


def compute_levels(nodes: list[dict[str, Any]], edges: list[dict[str, Any]]) -> dict[str, int]:
    explicit = {
        node["id"]: int(node["level"])
        for node in nodes
        if isinstance(node.get("level"), int) and node["level"] >= 0
    }
    levels = {node["id"]: explicit.get(node["id"], 0) for node in nodes}
    for _ in range(len(nodes)):
        changed = False
        for edge in edges:
            if edge["to"] in explicit:
                continue
            proposed = levels[edge["from"]] + 1
            if proposed > levels[edge["to"]]:
                levels[edge["to"]] = proposed
                changed = True
        if not changed:
            break
    minimum = min(levels.values(), default=0)
    return {node_id: level - minimum for node_id, level in levels.items()}


def node_dimensions(node: dict[str, Any]) -> tuple[float, float]:
    style = node.get("style")
    shape = node.get("shape")
    if style == "person" or shape == "ellipse":
        return number(node.get("width"), 1.20), number(node.get("height"), 1.20)
    if style == "target":
        return number(node.get("width"), 3.60), number(node.get("height"), 0.92)
    return number(node.get("width"), 3.05), number(node.get("height"), 0.92)


def add_arrow_end(connector: Any) -> None:
    line = connector._element.spPr.get_or_add_ln()
    for existing in line.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd"):
        line.remove(existing)
    tail = OxmlElement("a:tailEnd")
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    line.append(tail)


def set_east_asian_font(run: Any, typeface: str = FONT) -> None:
    """Set the DrawingML East Asian font, which python-pptx does not expose."""
    properties = run._r.get_or_add_rPr()
    properties.set("lang", "zh-CN")
    properties.set("altLang", "en-US")
    namespace = "{http://schemas.openxmlformats.org/drawingml/2006/main}ea"
    for existing in properties.findall(namespace):
        properties.remove(existing)
    east_asian = OxmlElement("a:ea")
    east_asian.set("typeface", typeface)
    properties.append(east_asian)


def style_line(connector: Any, *, width: float = 1.25, arrow: bool = False) -> None:
    connector.line.color.rgb = BLACK
    connector.line.width = Pt(width)
    connector.line.dash_style = MSO_LINE_DASH_STYLE.SOLID
    if arrow:
        add_arrow_end(connector)


def add_text_box(
    slide: Any,
    text: str,
    left: float,
    top: float,
    width: float,
    height: float,
    *,
    font_size: float,
    bold: bool = False,
    color: RGBColor = BLACK,
    fill: RGBColor | None = None,
    line: RGBColor | None = None,
    name: str | None = None,
) -> Any:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    if name:
        box.name = name
    if fill is not None:
        box.fill.solid()
        box.fill.fore_color.rgb = fill
    else:
        box.fill.background()
    if line is not None:
        box.line.color.rgb = line
        box.line.width = Pt(1)
    else:
        box.line.fill.background()
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.05)
    frame.margin_right = Inches(0.05)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT
    set_east_asian_font(run)
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_node(slide: Any, node: dict[str, Any], position: dict[str, float]) -> Any:
    style = node.get("style", "company")
    ellipse = node.get("shape") == "ellipse" or style == "person"
    geometry = MSO_AUTO_SHAPE_TYPE.OVAL if ellipse else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(
        geometry,
        Inches(position["x"]),
        Inches(position["y"]),
        Inches(position["w"]),
        Inches(position["h"]),
    )
    shape.name = f"node-{node['id']}"

    if style == "person":
        fill, border, text_color = ORANGE, ORANGE, WHITE
    elif style == "target":
        fill, border, text_color = DARK_BLUE, DARK_BLUE, WHITE
    else:
        fill, border, text_color = WHITE, BLUE, BLACK
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = border
    shape.line.width = Pt(1.25)

    label = str(node["label"])
    length = len(label.replace("\n", ""))
    if ellipse:
        font_size = 14.5 if length <= 4 else 12.5
    elif length <= 10:
        font_size = 17.0
    elif length <= 20:
        font_size = 14.0
    elif length <= 30:
        font_size = 12.5
    else:
        font_size = 11.5
    if style == "target":
        font_size = max(font_size, 16.0)
    font_size = number(node.get("font_size"), font_size)

    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.08)
    frame.margin_top = Inches(0.05)
    frame.margin_bottom = Inches(0.05)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = label
    run.font.name = FONT
    set_east_asian_font(run)
    run.font.size = Pt(font_size)
    run.font.color.rgb = text_color
    run.font.bold = style == "target"
    return shape


def find_cycles(node_ids: set[str], edges: list[dict[str, Any]]) -> bool:
    adjacency: dict[str, list[str]] = defaultdict(list)
    indegree = {node_id: 0 for node_id in node_ids}
    for edge in edges:
        adjacency[edge["from"]].append(edge["to"])
        indegree[edge["to"]] += 1
    stack = [node_id for node_id, degree in indegree.items() if degree == 0]
    visited = 0
    while stack:
        current = stack.pop()
        visited += 1
        for target in adjacency[current]:
            indegree[target] -= 1
            if indegree[target] == 0:
                stack.append(target)
    return visited != len(node_ids)


def compute_tree_layout(
    nodes: list[dict[str, Any]],
    edges: list[dict[str, Any]],
    levels: dict[str, int],
) -> tuple[dict[str, dict[str, float]], float, float]:
    node_by_id = {node["id"]: node for node in nodes}
    incoming: dict[str, list[str]] = defaultdict(list)
    outgoing: dict[str, list[str]] = defaultdict(list)
    for edge in edges:
        incoming[edge["to"]].append(edge["from"])
        outgoing[edge["from"]].append(edge["to"])
    for target, sources in incoming.items():
        sources.sort(key=lambda node_id: (number(node_by_id[node_id].get("order"), 9999), node_id))

    max_level = max(levels.values(), default=0)
    roots = [node["id"] for node in nodes if not outgoing[node["id"]]]
    roots.sort(key=lambda node_id: (number(node_by_id[node_id].get("order"), 9999), node_id))
    if not roots:
        raise ValueError("ownership graph has no final investee/root")

    sibling_gap = 0.30
    root_gap = 0.55
    memo: dict[str, float] = {}
    visiting: set[str] = set()

    def span(node_id: str) -> float:
        if node_id in memo:
            return memo[node_id]
        if node_id in visiting:
            raise ValueError("ownership graph contains a cycle")
        visiting.add(node_id)
        own_width, _ = node_dimensions(node_by_id[node_id])
        child_ids = incoming[node_id]
        child_width = sum(span(child) for child in child_ids)
        if child_ids:
            child_width += sibling_gap * (len(child_ids) - 1)
        memo[node_id] = max(own_width, child_width)
        visiting.remove(node_id)
        return memo[node_id]

    content_width = sum(span(root) for root in roots) + root_gap * (len(roots) - 1)
    max_content_width = 53.0
    horizontal_scale = min(1.0, max_content_width / max(content_width, 0.1))
    margin_x = 0.65
    slide_width = min(55.5, content_width * horizontal_scale + margin_x * 2)
    slide_width = max(slide_width, 13.333)
    offset = (slide_width - content_width * horizontal_scale) / 2
    level_top = 1.55
    level_gap = 2.45
    positions: dict[str, dict[str, float]] = {}

    def assign(node_id: str, left: float, allotted: float) -> None:
        node = node_by_id[node_id]
        width, height = node_dimensions(node)
        width *= horizontal_scale
        center = left + allotted / 2
        positions[node_id] = {
            "x": center - width / 2,
            "y": level_top + levels[node_id] * level_gap,
            "w": width,
            "h": height,
        }
        child_ids = incoming[node_id]
        if not child_ids:
            return
        child_total = sum(span(child) * horizontal_scale for child in child_ids)
        child_total += sibling_gap * horizontal_scale * (len(child_ids) - 1)
        cursor = left + (allotted - child_total) / 2
        for child in child_ids:
            child_span = span(child) * horizontal_scale
            assign(child, cursor, child_span)
            cursor += child_span + sibling_gap * horizontal_scale

    cursor = offset
    assigned: set[str] = set()
    for root in roots:
        allotted = span(root) * horizontal_scale
        assign(root, cursor, allotted)
        assigned.update(positions)
        cursor += allotted + root_gap * horizontal_scale

    # DAG nodes shared by branches should normally be duplicated in graph.json.
    # This fallback keeps any unassigned node visible instead of silently dropping it.
    missing = [node for node in nodes if node["id"] not in assigned]
    for index, node in enumerate(missing):
        width, height = node_dimensions(node)
        positions[node["id"]] = {
            "x": margin_x + index * (width + 0.3),
            "y": level_top + levels[node["id"]] * level_gap,
            "w": width,
            "h": height,
        }

    slide_height = max(7.5, level_top + max_level * level_gap + 1.65)
    return positions, slide_width, slide_height


def add_holding_label(
    slide: Any,
    edge: dict[str, Any],
    source: dict[str, float],
    target: dict[str, float],
) -> None:
    label = edge.get("holding") or edge.get("label")
    if not label:
        return
    width = 0.95 if len(str(label)) <= 8 else 1.35
    left = source["x"] + source["w"] / 2 - width / 2
    top = min(source["y"] + source["h"] + 0.10, target["y"] - 0.62)
    add_text_box(
        slide,
        str(label),
        left,
        top,
        width,
        0.32,
        font_size=11.5,
        fill=WHITE,
        name=f"holding-{edge['from']}-{edge['to']}",
    )


def render_connectors(
    slide: Any,
    edges: list[dict[str, Any]],
    positions: dict[str, dict[str, float]],
) -> None:
    by_target: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for edge in edges:
        by_target[edge["to"]].append(edge)

    for target_id, target_edges in by_target.items():
        target = positions[target_id]
        target_x = target["x"] + target["w"] / 2
        target_y = target["y"]
        target_edges.sort(key=lambda edge: positions[edge["from"]]["x"])

        if len(target_edges) == 1:
            edge = target_edges[0]
            source = positions[edge["from"]]
            source_x = source["x"] + source["w"] / 2
            source_y = source["y"] + source["h"]
            connector_type = (
                MSO_CONNECTOR.STRAIGHT
                if abs(source_x - target_x) < 0.03
                else MSO_CONNECTOR.ELBOW
            )
            connector = slide.shapes.add_connector(
                connector_type,
                Inches(source_x),
                Inches(source_y),
                Inches(target_x),
                Inches(target_y),
            )
            connector.name = f"edge-{edge['from']}-{edge['to']}"
            style_line(connector, arrow=True)
            add_holding_label(slide, edge, source, target)
            continue

        source_points = []
        for edge in target_edges:
            source = positions[edge["from"]]
            source_points.append((source["x"] + source["w"] / 2, source["y"] + source["h"], edge))
        highest_target_space = min(target_y - source_y for _, source_y, _ in source_points)
        bus_y = target_y - max(0.38, min(0.68, highest_target_space * 0.34))
        min_x = min(point[0] for point in source_points)
        max_x = max(point[0] for point in source_points)

        bus = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(min_x),
            Inches(bus_y),
            Inches(max_x),
            Inches(bus_y),
        )
        bus.name = f"bus-{target_id}"
        style_line(bus)
        for source_x, source_y, edge in source_points:
            stub = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                Inches(source_x),
                Inches(source_y),
                Inches(source_x),
                Inches(bus_y),
            )
            stub.name = f"edge-{edge['from']}-{edge['to']}"
            style_line(stub)
            add_holding_label(slide, edge, positions[edge["from"]], target)
        arrow = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(target_x),
            Inches(bus_y),
            Inches(target_x),
            Inches(target_y),
        )
        arrow.name = f"arrow-{target_id}"
        style_line(arrow, arrow=True)


def render(graph: dict[str, Any], output: Path) -> None:
    nodes = graph.get("nodes", [])
    edges = graph.get("edges", [])
    if not nodes:
        raise ValueError("graph must contain at least one node")
    node_ids = {node["id"] for node in nodes}
    if find_cycles(node_ids, edges):
        raise ValueError("ownership graph contains a cycle")
    levels = compute_levels(nodes, edges)
    positions, slide_width, slide_height = compute_tree_layout(nodes, edges, levels)

    presentation = Presentation()
    presentation.slide_width = Inches(slide_width)
    presentation.slide_height = Inches(slide_height)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = WHITE

    add_text_box(
        slide,
        str(graph.get("title") or "股权穿透图谱"),
        0.65,
        0.20,
        slide_width - 1.30,
        0.55,
        font_size=26,
        bold=True,
        name="chart-title",
    )

    # Relationships go behind nodes and labels.
    render_connectors(slide, edges, positions)
    for node in nodes:
        position = positions[node["id"]]
        add_node(slide, node, position)
        if node.get("tag"):
            text = str(node["tag"])
            is_badge = node.get("tag_style") == "badge" or "实际控制人" in text
            width = min(max(position["w"], 1.50), 2.25) if is_badge else position["w"]
            height = 0.55 if "\n" in text else 0.34
            add_text_box(
                slide,
                text,
                position["x"] + position["w"] / 2 - width / 2,
                position["y"] - height - 0.10,
                width,
                height,
                font_size=9.5 if "\n" in text else 11,
                bold=is_badge,
                color=WHITE if is_badge else BLACK,
                fill=RED if is_badge else None,
                line=RED if is_badge else None,
                name=f"tag-{node['id']}",
            )
        if node.get("collapsed"):
            size = 0.25
            marker = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.OVAL,
                Inches(position["x"] + position["w"] / 2 - size / 2),
                Inches(position["y"] - 0.40),
                Inches(size),
                Inches(size),
            )
            marker.name = f"collapsed-{node['id']}"
            marker.fill.solid()
            marker.fill.fore_color.rgb = BLUE
            marker.line.color.rgb = BLUE
            marker.text = "+"
            frame = marker.text_frame
            frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
            frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            paragraph = frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.CENTER
            run = paragraph.runs[0]
            run.font.name = FONT
            set_east_asian_font(run)
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = WHITE

    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    print(
        f"Created editable PowerPoint: {output} "
        f"({len(nodes)} nodes, {len(edges)} relationships, {max(levels.values()) + 1} levels)"
    )


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    try:
        graph = json.loads(args.input.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise SystemExit(f"Cannot read graph JSON: {exc}") from exc
    try:
        render(graph, args.output)
    except (KeyError, TypeError, ValueError) as exc:
        raise SystemExit(f"Cannot render graph: {exc}") from exc


if __name__ == "__main__":
    main()
